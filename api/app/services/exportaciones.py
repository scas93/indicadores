"""Fase 3 — Cuenta Pública (38 columnas) y Transparencia (26): misma tabla para la vista previa
(JSON) y para el .xlsx. Estructura de encabezados, orden y agrupación tomados de los exports reales
del sistema de referencia (humanos/docs/referencias/*.xls). Son de UN programa y un año.

Vacías a propósito (el sistema de referencia las deja vacías y aquí no hay el dato): ejercicio fiscal,
clasificación funcional, Aprobado…Pagado (Fase 4), meta modificada, fechas del periodo, sentido, metas
ajustadas, fecha de validación/actualización, nota.
De la ficha del indicador: ID del indicador = `id_numerico`; meta programada (texto) = Meta administración;
descripción de variables / método de cálculo = fórmula con las descripciones de A y B; "Fuente de
información" = medios de verificación (verificado idéntico en la referencia); "Valor línea base" =
resultado programado del año base (confirmado con la referencia).
PENDIENTE DE VERIFICACIÓN: "Definición del indicador" = Interpretación del resultado (hipótesis).
"""
import io
import uuid
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import calculo_estadisticas as ce
from ..deps import ActorCtx
from ..models import ClasificacionProgramatica, Frecuencia, MetaAnual
from . import estadisticas as est

CP_GRUPOS = [("", 2), ("Programa o proyecto de inversión", 5),
             ("Presupuesto del programa presupuestario", 5), ("MIR", 4), ("Indicadores", 7), ("", 15)]
CP_ENCABEZADOS = [
    "Año", "Ejercicio fiscal", "Clasificación del programa presupuestario",
    "Clave del programa presupuestario", "Nombre del programa presupuestario",
    "Clasificación funcional del gasto al que corresponde al programa presupuestario",
    "Nombre de la dependencia o entendidad que lo ejecuta", "Aprobado", "Modificado", "Devengado",
    "Ejercido", "Pagado", "Cuenta con MIR (SI/NO)", "Nivel de la MIR del programa",
    "Descripción del resumen narrativo (FIN, Propósito, Componentes y Actividades)", "ID del indicador",
    "Nombre del indicador", "Nivel de la MIR, al que corresponde el indicador", "Fórmula de cálculo",
    "Descripción de variables de la fórmula", "Meta del indicador programada",
    "Meta del indicador modficada", "Meta del indicador alcanzada", "Avance / Programado", "variable",
    "Unidad de medida", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
    "Septiembre", "Octubre", "Noviembre", "Diciembre"]
CP_COMBINADAS = 24  # las primeras 24 columnas se combinan en las 2 filas (A y B) de cada indicador

TR_ENCABEZADOS = [
    "Ejercicio", "Fecha de inicio del periodo que se informa", "Fecha de terminación del periodo que se informa",
    "Nombre del programa", "Objetivo institucional (Resumen narrativo)", "Nivel de la MIR", "ID del indicador",
    "Nombre del indicador", "Dimensión a medir", "Definición del indicador", "Algoritmo", "Método de cálculo",
    "Unidad de medida del indicador", "Frecuencia de medición", "Año línea base", "Valor línea base",
    "Metas programadas del año en análisis", "Descripción de la meta programada", "Metas ajustadas en su caso",
    "Avance de las metas", "Sentido del indicador", "Fuente de información", "Fecha de validación",
    "Área responsable de la información", "Fecha de actualización", "Nota"]
TR_PREVIAS = [["Título", "Nombre Corto", "Descripción"], ["", "", ""]]

_NIVEL = {"fin": "Fin", "proposito": "Proposito", "componente": "Componente", "actividad": "Actividad"}
_DIMENSION = {"eficacia": "Eficacia", "eficiencia": "Eficiencia", "calidad": "Calidad", "economia": "Economía"}
_FORMULA = {"a": "A", "a_sobre_b_pct": "(A / B) * 100", "a_sobre_b_menos1_pct": "((A / B) - 1) * 100",
            "a_sobre_b": "A / B"}


def _n(d: Decimal | float | None):
    """Número para la celda; vacío = ''. Los cálculos van a 2 decimales (la referencia no redondea)."""
    return "" if d is None else float(d)


def _formula(f: est.Fila) -> str:
    return _FORMULA.get(f.ind.algoritmo.value, "") if f.ind.algoritmo else ""


def _variables(f: est.Fila) -> str:
    """Fórmula con las descripciones de las variables (CP "Descripción de variables de la fórmula" y
    Transparencia "Método de cálculo"), formato de la referencia."""
    return ce.formula_con_descripciones(f.ind.algoritmo.value if f.ind.algoritmo else None,
                                        f.ind.descripcion_a, f.ind.descripcion_b)


def _filas_programa(db: Session, actor: ActorCtx, anio: int, centro_gestor_id, programa_id):
    p, cg = est.programa_exportable(db, actor, centro_gestor_id, programa_id)
    return p, cg, est.universo(db, actor, anio, programa_id=p.id)


def cuenta_publica(db: Session, actor: ActorCtx, anio: int, centro_gestor_id: uuid.UUID,
                   programa_id: uuid.UUID) -> dict:
    p, cg, filas = _filas_programa(db, actor, anio, centro_gestor_id, programa_id)
    clasif = db.get(ClasificacionProgramatica, p.clasificacion_programatica_id) \
        if p.clasificacion_programatica_id else None
    out = []
    for f in filas:
        valor_ce, _ = f.evaluar(ce.CO_CE)
        comunes = [str(anio), "", clasif.nombre if clasif else "", p.clave, p.nombre, "", cg.nombre,
                   "", "", "", "", "", "Si", _NIVEL[f.elem.nivel.value], f.elem.resumen_narrativo, f.ind.id_numerico,
                   f.ind.nombre, _NIVEL[f.elem.nivel.value], _formula(f), _variables(f), f.ind.meta_administracion or "", "",
                   _n(f.res.cumplimiento), _n(valor_ce)]
        usa_b = f.ind.algoritmo is not None and f.ind.algoritmo.value != "a"
        va = [f.avances.get(m, (None, None))[0] for m in range(1, 13)]
        vb = [f.avances.get(m, (None, None))[1] if usa_b else None for m in range(1, 13)]
        out.append(comunes + ["A", f.ind.unidad_a] + [_n(x) for x in va])
        out.append([None] * CP_COMBINADAS + ["B", (f.ind.unidad_b or "") if usa_b else ""] + [_n(x) for x in vb])
    return {"titulo": f"{cg.nombre}\nINDICADORES DE RESULTADOS", "grupos": CP_GRUPOS,
            "encabezados": CP_ENCABEZADOS, "filas": out, "filas_por_indicador": 2,
            "columnas_combinadas": CP_COMBINADAS, "nombre_archivo": f"cuenta_publica_{p.clave}_{anio}.xlsx"}


def transparencia(db: Session, actor: ActorCtx, anio: int, centro_gestor_id: uuid.UUID,
                  programa_id: uuid.UUID) -> dict:
    p, cg, filas = _filas_programa(db, actor, anio, centro_gestor_id, programa_id)
    out = []
    for f in filas:
        frec = db.get(Frecuencia, f.ind.frecuencia_id) if f.ind.frecuencia_id else None
        out.append([
            str(anio), "", "", p.nombre, f.elem.resumen_narrativo, _NIVEL[f.elem.nivel.value], f.ind.id_numerico,
            f.ind.nombre, _DIMENSION.get(f.ind.dimension.value, "") if f.ind.dimension else "",
            f.ind.interpretacion,  # "Definición del indicador" = Interpretación: PENDIENTE DE VERIFICACIÓN (hipótesis)
            _formula(f), _variables(f), f.ind.unidad_medida, frec.nombre if frec else "",
            f.ind.anio_base if f.ind.anio_base is not None else "", _n(_valor_linea_base(db, f)),
            _n(f.meta_num), f.ind.meta_administracion or "", "", _n(f.res.cumplimiento), "", f.elem.medios_verificacion, "",
            cg.nombre, "", ""])
    return {"previas": TR_PREVIAS, "encabezados": TR_ENCABEZADOS, "filas": out,
            "nombre_archivo": f"transparencia_{p.clave}_{anio}.xlsx"}


def _valor_linea_base(db: Session, f: est.Fila) -> Decimal | None:
    if f.ind.anio_base is None:
        return None
    m = db.scalar(select(MetaAnual).where(MetaAnual.indicador_id == f.ind.id,
                                          MetaAnual.anio == f.ind.anio_base))
    alg = f.ind.algoritmo.value if f.ind.algoritmo else None
    return ce.meta_numerica(alg, m.valor_a_programado, m.valor_b_programado) if m else None


# --------------------------------------------------------------------------- .xlsx
_GRIS = PatternFill("solid", fgColor="BBBBBB")
_BORDE = Border(*(Side(style="thin"),) * 4)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _estilo_encabezado(c) -> None:
    c.fill, c.font, c.alignment, c.border = _GRIS, Font(bold=True), _CENTRO, _BORDE


def _valor(v):
    return None if v == "" else v


def cuenta_publica_xlsx(t: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuenta Pública"
    n = len(t["encabezados"])
    ws.append([t["titulo"]] + [None] * (n - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.row_dimensions[1].height = 34
    _estilo_encabezado(ws.cell(1, 1))
    col = 1
    for texto, span in t["grupos"]:
        ws.cell(2, col, texto or None)
        if span > 1:
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + span - 1)
        for k in range(col, col + span):
            _estilo_encabezado(ws.cell(2, k))
        col += span
    ws.append(t["encabezados"])
    for c in ws[3]:
        _estilo_encabezado(c)
    ws.row_dimensions[3].height = 60
    for i, fila in enumerate(t["filas"]):
        ws.append([_valor(v) for v in fila])
        r = ws.max_row
        for c in ws[r]:
            c.border, c.alignment = _BORDE, Alignment(vertical="center", wrap_text=True)
        if i % t["filas_por_indicador"] == t["filas_por_indicador"] - 1:  # cierra el par A/B
            for k in range(1, t["columnas_combinadas"] + 1):
                ws.merge_cells(start_row=r - 1, start_column=k, end_row=r, end_column=k)
    _anchos(ws, n)
    ws.freeze_panes = "A4"
    return _bytes(wb)


def transparencia_xlsx(t: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transparencia"
    n = len(t["encabezados"])
    for previa in t["previas"]:
        ws.append(previa + [None] * (n - len(previa)))
        ws.merge_cells(start_row=ws.max_row, start_column=4, end_row=ws.max_row, end_column=n)
        for c in ws[ws.max_row][:3]:
            _estilo_encabezado(c)
    ws.append(t["encabezados"])
    for c in ws[ws.max_row]:
        _estilo_encabezado(c)
    ws.row_dimensions[ws.max_row].height = 45
    for fila in t["filas"]:
        ws.append([_valor(v) for v in fila])
        for c in ws[ws.max_row]:
            c.border, c.alignment = _BORDE, Alignment(vertical="center", wrap_text=True)
    _anchos(ws, n)
    ws.freeze_panes = ws.cell(len(t["previas"]) + 2, 1)
    return _bytes(wb)


def _anchos(ws, n: int) -> None:
    for i in range(1, n + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
